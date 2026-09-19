from datetime import date, datetime, timezone
from decimal import Decimal

from app.ingest.commit import commit_snapshot
from app.ingest.parser import ParsedHolding
from app.models import NavPoint, PriceHistoryCache, QuoteCache
from app.performance import (
    add_nav_point,
    build_performance,
    delete_nav_point,
    list_nav_points,
    parse_nav_csv,
    _blend_series,
    _value_at,
)
from sqlalchemy import select


def _h(t, u, c, isin=None, cash=False):
    return ParsedHolding(raw_ticker=f"{t} US", ticker=t, exchange="US",
                         units=Decimal(str(u)), avg_cost=Decimal(str(c)), isin=isin, is_cash=cash)


def _seed_price_series(session, ticker, pairs):
    session.add(PriceHistoryCache(ticker=ticker, series=[[d, str(c)] for d, c in pairs],
                                  fetched_at=datetime.now(timezone.utc)))


def _seed_quote(session, ticker, price):
    session.add(QuoteCache(ticker=ticker, fmp_symbol=ticker, price=Decimal(str(price)),
                           prev_close=Decimal(str(price)), ok=True, fetched_at=datetime.now(timezone.utc)))


# ---- NAV CRUD ----

def test_nav_add_list_delete_upsert(session):
    add_nav_point(session, date(2026, 7, 16), Decimal("1646.90"))
    add_nav_point(session, date(2026, 6, 16), Decimal("1500"))
    pts = list_nav_points(session)
    assert [p.as_of for p in pts] == [date(2026, 6, 16), date(2026, 7, 16)]   # sorted

    # upsert same date
    add_nav_point(session, date(2026, 7, 16), Decimal("1650"))
    assert len(list_nav_points(session)) == 2
    assert session.execute(select(NavPoint).where(NavPoint.as_of == date(2026, 7, 16))).scalar_one().nav_per_share == Decimal("1650")

    delete_nav_point(session, date(2026, 6, 16))
    assert len(list_nav_points(session)) == 1


def test_parse_nav_csv():
    csv = b"date,nav_per_share\n2026-06-16,1500.00\n2026-07-16,1646.90\n"
    pts = parse_nav_csv(csv)
    assert pts == [(date(2026, 6, 16), Decimal("1500.00")), (date(2026, 7, 16), Decimal("1646.90"))]
    # synonyms + thousands separators
    assert parse_nav_csv(b"as_of,price\n2026-07-16,\"1,646.90\"\n") == [(date(2026, 7, 16), Decimal("1646.90"))]


def test_parse_nav_xlsx_daily_nav_sheet():
    import io
    from datetime import datetime as dt
    from openpyxl import Workbook
    from app.performance import parse_nav_file

    wb = Workbook()
    wb.active.title = "Summary"          # active sheet is NOT the NAV one
    ws = wb.create_sheet("Daily NAV")
    ws.append(["Date", "NAV (USD)", "Daily Return"])
    ws.append([dt(2023, 11, 2), 1000.61, None])
    ws.append([dt(2026, 7, 16), 1646.90, -0.0348])
    buf = io.BytesIO(); wb.save(buf)

    pts = parse_nav_file("NextGen_NAV.xlsx", buf.getvalue())
    assert pts == [(date(2023, 11, 2), Decimal("1000.61")), (date(2026, 7, 16), Decimal("1646.90"))]


def test_parse_nav_xlsx_price_column():
    # Regression: a broker export with "Price_USD" instead of a NAV-named column
    # used to import 0 rows.
    import io
    from datetime import datetime as dt
    from openpyxl import Workbook
    from app.performance import parse_nav_file

    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Price_USD"])
    ws.append([dt(2026, 7, 31), 1583.38])
    ws.append([dt(2026, 8, 3), 1639.79])
    buf = io.BytesIO(); wb.save(buf)

    pts = parse_nav_file("price.xlsx", buf.getvalue())
    assert pts == [(date(2026, 7, 31), Decimal("1583.38")), (date(2026, 8, 3), Decimal("1639.79"))]


def test_parse_nav_csv_price_variant_header():
    assert parse_nav_csv(b"date,price_usd\n2026-07-31,1583.38\n") == \
        [(date(2026, 7, 31), Decimal("1583.38"))]


def test_bulk_add_nav_points_upserts(session):
    from app.performance import bulk_add_nav_points
    n = bulk_add_nav_points(session, [(date(2023, 11, 2), Decimal("1000.61")),
                                      (date(2026, 7, 16), Decimal("1646.90"))])
    assert n == 2 and len(list_nav_points(session)) == 2
    # re-import with a changed value updates in place, no duplicate
    bulk_add_nav_points(session, [(date(2026, 7, 16), Decimal("1650"))])
    assert len(list_nav_points(session)) == 2
    assert list_nav_points(session)[-1].nav_per_share == Decimal("1650")


# ---- benchmark helpers ----

def test_value_at_modes():
    s = [("2026-06-01", Decimal("10")), ("2026-06-15", Decimal("12")), ("2026-06-30", Decimal("15"))]
    assert _value_at(s, "2026-06-10", mode="after") == Decimal("12")
    assert _value_at(s, "2026-06-20", mode="before") == Decimal("12")


def test_blend_is_5050():
    qqq = [("d1", Decimal("100")), ("d2", Decimal("110"))]   # +10%
    smh = [("d1", Decimal("50")), ("d2", Decimal("60"))]      # +20%
    b = _blend_series(qqq, smh)
    # rebased: d1 -> 1.0 ; d2 -> 0.5*1.1 + 0.5*1.2 = 1.15
    assert b[0][1] == Decimal("1.0")
    assert b[1][1] == Decimal("1.15")


# ---- fund + attribution ----

def test_fund_return_needs_two_points(session):
    _seed_price_series(session, "QQQ", [("2026-07-16", Decimal("500")), ("2026-07-17", Decimal("505"))])
    _seed_price_series(session, "SPY", [("2026-07-16", Decimal("600")), ("2026-07-17", Decimal("606"))])
    add_nav_point(session, date(2026, 7, 16), Decimal("1000"))
    v = build_performance(session)
    assert v.has_fund_return is False
    assert v.fund_return is None

    add_nav_point(session, date(2026, 7, 17), Decimal("1100"))
    v = build_performance(session)
    assert v.has_fund_return is True
    assert v.fund_return == Decimal("0.1")            # 1100/1000 - 1
    # QQQ 505/500 - 1 = 1%
    assert round(v.returns["QQQ"], 4) == Decimal("0.01")


def test_contributions_sum_to_fund_return(session):
    commit_snapshot(holdings=[
        _h("AAA", 10, 10, isin="US0000000001"),    # cost 100 -> mv 200 : +100
        _h("BBB", 10, 10, isin="US0000000002"),    # cost 100 -> mv 50  : -50
        _h("USD Cash", 100, 1, cash=True),
    ], as_of=date(2026, 7, 1), filename="f", notes=None, pillar_assignments=None, session=session)
    _seed_quote(session, "AAA", 20)
    _seed_quote(session, "BBB", 5)
    session.commit()

    v = build_performance(session)
    # total cost 200, total mv 250 -> fund unrealized return = 25%
    assert v.fund_unrl_return == Decimal("0.25")
    total_contrib = sum(c.contribution for c in v.contributors)
    assert total_contrib == v.fund_unrl_return
    # AAA contributes +100/200 = +50%, BBB -50/200 = -25%
    by = {c.ticker: c.contribution for c in v.contributors}
    assert by["AAA"] == Decimal("0.5")
    assert by["BBB"] == Decimal("-0.25")


def test_nav_kpis():
    from app.models import NavPoint
    from app.performance import compute_nav_kpis

    def np(d, nav):
        return NavPoint(as_of=d, nav_per_share=Decimal(str(nav)))

    navs = [
        np(date(2025, 12, 31), 100),   # prior year-end / prior quarter-end base
        np(date(2026, 1, 31), 110),    # Jan +10%
        np(date(2026, 2, 27), 99),     # Feb -10% ; latest
    ]
    # QQQ reference series across the same window
    ref = [("2025-12-31", Decimal("500")), ("2026-01-31", Decimal("525")),
           ("2026-02-27", Decimal("500"))]
    k = compute_nav_kpis(navs, ref=ref, ref_label="QQQ")
    assert k["year"] == 2026 and k["ref_label"] == "QQQ"
    assert k["ytd"]["fund"] == Decimal("99") / Decimal("100") - 1        # -1%
    assert k["mtd"]["fund"] == Decimal("99") / Decimal("110") - 1        # Feb -10%
    assert k["qtd"]["fund"] == Decimal("99") / Decimal("100") - 1        # Q1 vs 2025 year-end
    # QQQ YTD over 2025-12-31 -> 2026-02-27: 500/500 - 1 = 0
    assert k["ytd"]["ref"] == Decimal("0")
    assert [lbl for lbl, _, _ in k["by_month"]] == ["Jan", "Feb"]
    assert k["by_month"][0][1] == Decimal("110") / Decimal("100") - 1    # fund Jan +10%
    assert k["by_month"][0][2] == Decimal("525") / Decimal("500") - 1    # QQQ Jan +5%
    assert k["by_quarter"][0][:2] == ("Q1", Decimal("99") / Decimal("100") - 1)
    # prior full years only (2026 is the current YTD year, excluded); 2025 is the
    # inception year here so it is measured from the first NAV point.
    assert k["by_year"] == [("2025", Decimal("0"), Decimal("0"), True)]


def test_nav_kpis_annual_full_years():
    from app.models import NavPoint
    from app.performance import compute_nav_kpis

    def np(d, nav):
        return NavPoint(as_of=d, nav_per_share=Decimal(str(nav)))

    navs = [np(date(2023, 11, 2), 1000), np(date(2023, 12, 31), 1100),   # 2023 incep +10%
            np(date(2024, 12, 31), 1320),                                # 2024 +20%
            np(date(2025, 6, 30), 1400)]                                 # current year (2025)
    k = compute_nav_kpis(navs, ref=None)
    years = {label: (fr, incep) for label, fr, rr, incep in k["by_year"]}
    assert years["2023"] == (Decimal("1100") / Decimal("1000") - 1, True)     # from inception
    assert years["2024"] == (Decimal("1320") / Decimal("1100") - 1, False)    # full year +20%
    assert "2025" not in years                                                 # current year excluded





def test_nav_kpis_needs_two_points():
    from app.models import NavPoint
    from app.performance import compute_nav_kpis
    assert compute_nav_kpis([NavPoint(as_of=date(2026, 7, 16), nav_per_share=Decimal("1000"))], ref=[]) == {}


def test_pillar_contribution(session):
    commit_snapshot(holdings=[
        _h("AAA", 10, 10, isin="US0000000001"),
        _h("BBB", 10, 10, isin="US0000000002"),
    ], as_of=date(2026, 7, 1), filename="f", notes=None, pillar_assignments={"AAA": "Semis", "BBB": "Semis"},
        session=session)
    _seed_quote(session, "AAA", 20)
    _seed_quote(session, "BBB", 20)
    session.commit()
    v = build_performance(session)
    semis = next(p for p in v.by_pillar if p["pillar"] == "Semis")
    assert semis["contribution"] == Decimal("1")     # both doubled: (200)/(200) cost... +100/100? check


# ---- pillar vs benchmark ETF ----

def _seed_pillar_book(session):
    from app.models import Pillar, Security
    session.add_all([
        Pillar(id="P01", name="Semis", primary_etf="SMH"),
        Pillar(id="P03", name="Neoclouds", primary_etf=None, alt_etf="SKYY", caveat="weak proxy"),
    ])
    session.flush()
    commit_snapshot(holdings=[_h("AAA", 10, 10), _h("BBB", 20, 10), _h("CCC", 5, 10)],
                    as_of=date(2026, 7, 1), filename="f", notes=None,
                    pillar_assignments={"AAA": "Semis", "BBB": "Semis", "CCC": "Neoclouds"},
                    session=session)
    for t in ("AAA", "BBB", "CCC"):
        _seed_quote(session, t, 20)
    # Prior-year close -> latest close. AAA 10->20, BBB 10->15, SMH 100->130, SKYY 50->40.
    _seed_price_series(session, "AAA", [("2025-12-31", 10), ("2026-09-18", 20)])
    _seed_price_series(session, "BBB", [("2025-12-31", 10), ("2026-09-18", 15)])
    _seed_price_series(session, "CCC", [("2026-03-01", 10), ("2026-09-18", 12)])   # starts too late for YTD
    _seed_price_series(session, "SMH", [("2025-12-30", 90), ("2025-12-31", 100), ("2026-09-18", 130)])
    _seed_price_series(session, "SKYY", [("2025-12-31", 50), ("2026-09-18", 40)])
    session.commit()


def test_pillar_vs_etf_same_window(session):
    _seed_pillar_book(session)
    v = build_performance(session, pillar_period="ytd", today=date(2026, 9, 19))
    rows = {r["pillar"]: r for r in v.by_pillar}

    semis = rows["Semis"]
    # holdings-weighted: (10*20 + 20*15) / (10*10 + 20*10) - 1 = 500/300 - 1
    assert semis["pillar_return"] == Decimal(500) / Decimal(300) - 1
    assert semis["etf"] == "SMH" and semis["etf_return"] == Decimal("0.3")
    assert semis["excess"] == semis["pillar_return"] - Decimal("0.3")
    assert (semis["covered"], semis["total"]) == (2, 2)

    neo = rows["Neoclouds"]
    assert neo["etf"] == "SKYY" and neo["etf_is_alt"] and neo["caveat"] == "weak proxy"
    assert neo["etf_return"] == Decimal("-0.2")
    assert neo["pillar_return"] is None and neo["excess"] is None   # no YTD history
    assert (neo["covered"], neo["total"]) == (0, 1)
    assert v.pillar_period_end == "2026-09-18"


def test_pillar_period_windows(session):
    _seed_pillar_book(session)
    v = build_performance(session, pillar_period="3m", today=date(2026, 9, 19))
    neo = {r["pillar"]: r for r in v.by_pillar}["Neoclouds"]
    assert neo["pillar_return"] == Decimal("0.2")      # CCC history now spans the window
    assert build_performance(session, pillar_period="bogus").pillar_period == "ytd"


def test_pillar_etfs_included_in_refresh(session):
    from app.performance import pillar_etf_symbols
    _seed_pillar_book(session)
    assert pillar_etf_symbols(session) == ["SKYY", "SMH"]
