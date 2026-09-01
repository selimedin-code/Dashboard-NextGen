"""Parse a custodian holdings file into ParsedHolding rows.

Two formats, per the project decision:
  - Excel (.xlsx): the primary, reliable path. Columns are matched by header name,
    tolerant of ordering and extra columns. This is the format to standardize on.
  - PDF: a fallback for the custodian's rendered report. Parsed by x-position
    column bucketing (the layout export splits numbers on kerning, so token
    counting is unreliable — bucketing each fragment by its column is not).

Only four things are load-bearing: ticker, units, avg_cost, and (for identity)
ISIN/name. Everything else in the file is derived by the dashboard and ignored.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal

from app.ingest.normalize import (
    CASH_TICKER,
    ParseError,
    normalize_ticker,
    parse_decimal,
    validate_isin,
)


@dataclass
class ParsedHolding:
    raw_ticker: str          # as it appeared, e.g. "NVDA US"
    ticker: str              # normalized, no suffix, e.g. "NVDA"
    exchange: str | None     # "US", "GR", ... or None
    units: Decimal
    avg_cost: Decimal
    isin: str | None = None
    name: str | None = None
    is_cash: bool = False


@dataclass
class ParseResult:
    holdings: list[ParsedHolding]
    source_format: str                       # "excel" | "pdf"
    warnings: list[str] = field(default_factory=list)


class UnsupportedFileError(ValueError):
    pass


def parse_upload(filename: str, data: bytes) -> ParseResult:
    """Dispatch on file extension."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return parse_excel(data)
    if lower.endswith(".pdf"):
        return parse_pdf(data)
    raise UnsupportedFileError(
        f"Unsupported file type: {filename!r}. Upload a .xlsx or .pdf custodian file."
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

# Header synonyms -> canonical field. Matched case-insensitively on a normalized
# (alphanumeric-only) form so "Avg Cost", "avg_cost", "Average Cost" all map.
_HEADER_MAP = {
    "ticker": "ticker",
    "symbol": "ticker",
    "bbticker": "ticker",
    "bloomberg": "ticker",
    "units": "units",
    "totalunits": "units",
    "quantity": "units",
    "qty": "units",
    "shares": "units",
    "avgcost": "avg_cost",
    "averagecost": "avg_cost",
    "cost": "avg_cost",
    "costbasis": "avg_cost",
    "blendedcost": "avg_cost",
    "isin": "isin",
    "name": "name",
    "securityname": "name",
    "description": "name",
}


def _norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def parse_excel(data: bytes) -> ParseResult:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    # Find the header row: the first row that maps at least ticker + units + cost.
    header_index: dict[int, str] | None = None
    assumed_cost_col = False
    for raw_row in rows:
        mapping: dict[int, str] = {}
        for i, cell in enumerate(raw_row):
            field_name = _HEADER_MAP.get(_norm_header(cell))
            if field_name and field_name not in mapping.values():
                mapping[i] = field_name
        if {"ticker", "units", "avg_cost"}.issubset(mapping.values()):
            header_index = mapping
            break
        # Custodian variant: the avg-cost column exists but its header cell is
        # blank (e.g. "Ticker | ISIN | Name | Sector | Total units | <blank>").
        # If the column right of Units is unlabeled, assume it is Avg Cost and
        # say so in a warning the preview shows.
        if {"ticker", "units"}.issubset(mapping.values()):
            units_i = next(i for i, f in mapping.items() if f == "units")
            nxt = units_i + 1
            if nxt not in mapping and (nxt >= len(raw_row) or _norm_header(raw_row[nxt]) == ""):
                mapping[nxt] = "avg_cost"
                header_index = mapping
                assumed_cost_col = True
                break

    if header_index is None:
        raise ParseError(
            "Could not find a header row with Ticker, Units and Cost columns. "
            "Expected columns like: Ticker | ISIN | Name | Units | AvgCost."
        )

    holdings: list[ParsedHolding] = []
    warnings: list[str] = []
    if assumed_cost_col:
        warnings.append(
            "No Avg Cost header found — assumed the unlabeled column right of the "
            "Units column holds average cost. Check the preview values before committing."
        )
    for raw_row in rows:  # continues after the header row
        record = {field_name: raw_row[i] for i, field_name in header_index.items() if i < len(raw_row)}
        raw_ticker = record.get("ticker")
        if raw_ticker is None or str(raw_ticker).strip() == "":
            continue  # blank/spacer row
        holding = _build_holding(
            raw_ticker=str(raw_ticker).strip(),
            units_raw=record.get("units"),
            cost_raw=record.get("avg_cost"),
            isin_raw=record.get("isin"),
            name_raw=record.get("name"),
            warnings=warnings,
        )
        if holding is not None:
            holdings.append(holding)

    wb.close()
    return ParseResult(holdings=holdings, source_format="excel", warnings=warnings)


# ---------------------------------------------------------------------------
# PDF (custodian rendered report)
# ---------------------------------------------------------------------------

# x0 column boundaries (PDF points) for this custodian's layout. Fragments are
# bucketed by their left edge; everything left of `name` is the ticker/exch/ISIN
# block. Columns past `cost` are derived data and ignored.
_PDF_COLS = {
    "idblock": (0.0, 121.0),
    "name": (121.0, 207.0),
    "sector": (207.0, 268.0),
    "units": (268.0, 306.0),
    "cost": (306.0, 333.0),
}
# Anything right of this x0 is the sector-allocation summary block, not holdings.
_PDF_RIGHT_MARGIN = 560.0
_TOP_TOLERANCE = 3.0  # points; group word fragments into the same visual row


def _pdf_bucket(x0: float) -> str | None:
    for key, (lo, hi) in _PDF_COLS.items():
        if lo <= x0 < hi:
            return key
    return None


def parse_pdf(data: bytes) -> ParseResult:
    import pdfplumber

    holdings: list[ParsedHolding] = []
    warnings: list[str] = []
    seen_cash = False

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            words = [w for w in page.extract_words() if w["x0"] < _PDF_RIGHT_MARGIN]
            for frags in _group_rows(words):
                cells: dict[str, list[str]] = {k: [] for k in _PDF_COLS}
                for w in frags:
                    bucket = _pdf_bucket(w["x0"])
                    if bucket:
                        cells[bucket].append(w["text"])
                idblock = cells["idblock"]
                if not idblock:
                    continue
                symbol = idblock[0]

                # Cash line: "USD" in the id block, avg_cost is 1 by definition.
                if symbol.upper() == "USD":
                    if cells["units"]:
                        try:
                            units = parse_decimal("".join(cells["units"]))
                        except ParseError:
                            continue
                        holdings.append(
                            ParsedHolding(
                                raw_ticker=CASH_TICKER, ticker=CASH_TICKER, exchange=None,
                                units=units, avg_cost=Decimal("1"), is_cash=True,
                            )
                        )
                        seen_cash = True
                    continue

                idrest = "".join(idblock[1:])
                # idrest is exchange(2) + ISIN(12), possibly split across fragments.
                exchange = idrest[:2] if len(idrest) >= 2 else None
                isin = validate_isin(idrest[2:14]) if len(idrest) >= 14 else None
                if isin is None:
                    continue  # not a holdings row (header, footer, total)

                try:
                    units = parse_decimal("".join(cells["units"]))
                    cost = parse_decimal("".join(cells["cost"]))
                except ParseError:
                    continue

                raw_ticker = f"{symbol} {exchange}" if exchange else symbol
                holdings.append(
                    ParsedHolding(
                        raw_ticker=raw_ticker,
                        ticker=symbol,
                        exchange=exchange,
                        units=units,
                        avg_cost=cost,
                        isin=isin,
                        name=" ".join(cells["name"]) or None,
                    )
                )

    if not holdings:
        raise ParseError("No holdings rows found in the PDF. The layout may differ from the expected custodian report.")
    if not seen_cash:
        warnings.append("No 'USD Cash' line found in the PDF — verify the cash position was captured.")
    return ParseResult(holdings=holdings, source_format="pdf", warnings=warnings)


def _group_rows(words: list[dict]) -> list[list[dict]]:
    """Cluster word fragments into visual rows by their vertical position."""
    rows: dict[float, list[dict]] = {}
    for w in words:
        for top in rows:
            if abs(top - w["top"]) <= _TOP_TOLERANCE:
                rows[top].append(w)
                break
        else:
            rows[w["top"]] = [w]
    ordered = []
    for top in sorted(rows):
        ordered.append(sorted(rows[top], key=lambda w: w["x0"]))
    return ordered


# ---------------------------------------------------------------------------
# Shared
# ---------------------------------------------------------------------------


def _build_holding(
    *, raw_ticker: str, units_raw: object, cost_raw: object,
    isin_raw: object, name_raw: object, warnings: list[str],
) -> ParsedHolding | None:
    ticker, exchange = normalize_ticker(raw_ticker)
    is_cash = ticker == CASH_TICKER

    try:
        units = parse_decimal(units_raw)
    except ParseError:
        warnings.append(f"{raw_ticker}: could not read units ({units_raw!r}); row skipped.")
        return None

    if is_cash:
        avg_cost = Decimal("1")
    else:
        try:
            avg_cost = parse_decimal(cost_raw)
        except ParseError:
            warnings.append(f"{raw_ticker}: could not read cost ({cost_raw!r}); row skipped.")
            return None

    return ParsedHolding(
        raw_ticker=raw_ticker,
        ticker=ticker,
        exchange=exchange,
        units=units,
        avg_cost=avg_cost,
        isin=validate_isin(str(isin_raw) if isin_raw else None),
        name=(str(name_raw).strip() if name_raw else None),
        is_cash=is_cash,
    )
