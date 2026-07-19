"""Bulk pillar assignment — the 11-pillar thematic framework.

Two ways in, both landing in `securities.pillar`:
  - inline editing of currently-held tickers on the Pillars page
  - a bulk file (Excel/CSV of Ticker, Pillar) via "Upload Pillars"

A mapping may name tickers not yet held; those get a placeholder `securities`
row so the pillar is already in place when the position first appears.
"""

from __future__ import annotations

import csv
import io
import re

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ingest.normalize import normalize_ticker
from app.models import Security


class PillarParseError(ValueError):
    pass


def _norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


_TICKER_HEADERS = {"ticker", "symbol", "bbticker"}
_PILLAR_HEADERS = {"pillar", "theme", "bucket", "category", "sleeve"}


def parse_pillar_file(filename: str, data: bytes) -> dict[str, str]:
    """Parse a Ticker→Pillar mapping from Excel or CSV. Tickers are normalized
    (country suffix stripped) to match `securities.ticker`."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_excel(data)
    elif lower.endswith(".csv"):
        rows = _rows_from_csv(data)
    else:
        raise PillarParseError(f"Unsupported file: {filename!r}. Upload a .xlsx or .csv with Ticker and Pillar columns.")

    if not rows:
        raise PillarParseError("The file has no rows.")

    header = rows[0]
    t_idx = p_idx = None
    for i, cell in enumerate(header):
        h = _norm_header(cell)
        if t_idx is None and h in _TICKER_HEADERS:
            t_idx = i
        if p_idx is None and h in _PILLAR_HEADERS:
            p_idx = i
    if t_idx is None or p_idx is None:
        raise PillarParseError("Could not find a Ticker column and a Pillar column in the header row.")

    mapping: dict[str, str] = {}
    for row in rows[1:]:
        if t_idx >= len(row) or p_idx >= len(row):
            continue
        raw_ticker, pillar = row[t_idx], row[p_idx]
        if raw_ticker is None or str(raw_ticker).strip() == "":
            continue
        if pillar is None or str(pillar).strip() == "":
            continue
        ticker, _ = normalize_ticker(str(raw_ticker).strip())
        mapping[ticker] = str(pillar).strip()
    if not mapping:
        raise PillarParseError("No ticker/pillar pairs found in the file.")
    return mapping


def _rows_from_excel(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_from_csv(data: bytes) -> list[list]:
    text = data.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def apply_pillars(session: Session, mapping: dict[str, str], *, create_missing: bool = True) -> dict[str, int]:
    """Apply a ticker→pillar mapping. Returns {updated, created}. Caller commits."""
    updated = created = 0
    existing = {
        s.ticker: s for s in session.execute(select(Security)).scalars().all()
    }
    for ticker, pillar in mapping.items():
        sec = existing.get(ticker)
        if sec is None:
            if not create_missing:
                continue
            sec = Security(ticker=ticker, pillar=pillar, active=False)
            session.add(sec)
            created += 1
        else:
            sec.pillar = pillar
            updated += 1
    return {"updated": updated, "created": created}
